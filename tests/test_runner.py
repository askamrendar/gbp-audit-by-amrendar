import sys,pathlib,threading,unittest,urllib.request,urllib.error,json,tempfile
from unittest.mock import patch
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parents[1]))
import runner
from report import export_ppt,website_facts,FORM
class RunnerTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.server=runner.ThreadingHTTPServer(('127.0.0.1',0),runner.Handler);runner.PORT=cls.server.server_port
  cls.url=f'http://127.0.0.1:{runner.PORT}';threading.Thread(target=cls.server.serve_forever,daemon=True).start()
 @classmethod
 def tearDownClass(cls):cls.server.shutdown();cls.server.server_close()
 def request(self,path,token=True,origin=None,data=None):
  headers={'Origin':origin or self.url}
  if token:headers['Authorization']='Bearer '+runner.TOKEN
  if data is not None:headers['Content-Type']='application/json';data=json.dumps(data).encode()
  return urllib.request.urlopen(urllib.request.Request(self.url+path,headers=headers,data=data))
 def test_health_requires_token(self):
  with self.assertRaises(urllib.error.HTTPError) as c:self.request('/health',False)
  self.assertEqual(c.exception.code,403)
 def test_origin_rejected(self):
  with self.assertRaises(urllib.error.HTTPError):self.request('/health',origin='https://evil.test')
 def test_paired_health(self):self.assertTrue(json.load(self.request('/health'))['ready'])
 def test_static_interface(self):self.assertIn(b'GBP Audit by Amrendar',urllib.request.urlopen(self.url).read())
 def test_consent_required(self):
  with self.assertRaises(urllib.error.HTTPError) as c:self.request('/jobs',data={'consent':False})
  self.assertEqual(c.exception.code,400)
 def test_coordinate_pair(self):
  with self.assertRaises(ValueError):runner.validate_rows([{'store_code':'A','maps_url':'https://www.google.com/maps?cid=123','keyword':'shoe','search_latitude':20}])
 def test_schema_inventory(self):
  w=website_facts({'schema':[{'syntax':'parsed','value':{'@graph':[{'@type':'Store'}]}}]});self.assertEqual(w['Schema types'],'Store')
 def test_report_is_editable_and_complete(self):
  from pptx import Presentation
  with tempfile.TemporaryDirectory() as d:
   p=pathlib.Path(d)/'test.pptx';export_ppt({'input':{'store_code':'A','keyword':'shoe'},'started_at':'2026-10-22','profiles':[{'name':'Example Store','checks':[]}],'ranking':{}},p)
   r=Presentation(p);self.assertEqual(len(r.slides),7);self.assertEqual(len(FORM),21);self.assertGreaterEqual(sum(x.has_table for s in r.slides for x in s.shapes),5)
 def test_job_to_download(self):
  with tempfile.TemporaryDirectory() as d,patch.object(runner,'ROOT',pathlib.Path(d)),patch.object(runner.audit,'run') as run:
   def fake(rows,out,*a,**kw):out.mkdir(parents=True);(out/'audit-draft.pptx').write_bytes(b'test fixture')
   run.side_effect=fake
   response=json.load(self.request('/jobs',data={'consent':True,'rows':[{'store_code':'A','maps_url':'https://www.google.com/maps?cid=123','keyword':'shoe'}]}));jid=response['id']
   import time
   for _ in range(100):
    job=json.load(self.request('/jobs/'+jid))
    if job['status'] in ['completed','failed']:break
    time.sleep(.01)
   self.assertEqual(job['status'],'completed');self.assertIn('audit-results.zip',job['files']);self.assertEqual(self.request('/jobs/'+jid+'/file/audit-draft.pptx').read(),b'test fixture')
   with self.assertRaises(urllib.error.HTTPError):self.request('/jobs/'+jid+'/file/..%2FREADME.md')
 def test_invalid_logo_rejected(self):
  import base64
  with self.assertRaises(urllib.error.HTTPError) as c:
   self.request('/jobs',data={'consent':True,'rows':[{'store_code':'A','maps_url':'https://www.google.com/maps?cid=123','keyword':'shoe'}],'logo':base64.b64encode(b'not an image').decode()})
  self.assertEqual(c.exception.code,400)
 def test_logo_kept_in_local_job(self):
  import base64,io
  from PIL import Image
  raw=io.BytesIO();Image.new('RGB',(20,10),'blue').save(raw,'PNG')
  with tempfile.TemporaryDirectory() as tmp,patch.object(runner,'ROOT',pathlib.Path(tmp)),patch.object(runner.POOL,'submit') as submit:
   result=json.load(self.request('/jobs',data={'consent':True,'rows':[{'store_code':'A','maps_url':'https://www.google.com/maps?cid=123','keyword':'shoe'}],'brand':{'agency':'Test Agency','accent':'#123456'},'logo':base64.b64encode(raw.getvalue()).decode()}))
   brand=submit.call_args.args[-1];self.assertEqual(brand['agency'],'Test Agency');self.assertTrue(pathlib.Path(brand['logo_path']).is_file());self.assertIn(result['id'],brand['logo_path'])
   runner.JOBS.pop(result['id'])

if __name__=='__main__':unittest.main()

class ScreenshotReportTests(unittest.TestCase):
 def test_report_embeds_both_screenshots_and_all_form_criteria(self):
  from PIL import Image
  from pptx import Presentation
  from pptx.enum.shapes import MSO_SHAPE_TYPE
  with tempfile.TemporaryDirectory() as tmp:
   d=pathlib.Path(tmp);image=d/'gbp.png';Image.new('RGB',(1280,900),'white').save(image)
   data={'input':{'store_code':'A','keyword':'shoe'},'started_at':'2026-10-22','profiles':[{'name':'Example Store','screenshot':str(image),'checks':[]}],'ranking':{'screenshot':str(image),'requested_center':'28,77'},'search_location':{'source':'selected_profile_pin'}}
   export_ppt(data,d/'audit.pptx');r=Presentation(d/'audit.pptx')
   self.assertEqual(sum(x.shape_type==MSO_SHAPE_TYPE.PICTURE for x in r.slides[2].shapes),2)
   table_text=' '.join(c.text for s in r.slides for shape in s.shapes if shape.has_table for row in shape.table.rows for c in row.cells)
   notes=' '.join(slide.notes_slide.notes_text_frame.text for slide in r.slides)
   for criterion in FORM:self.assertIn(criterion,notes)
   self.assertIn('Photos / media shown',table_text)
   self.assertIn('Days 22–30',table_text)
