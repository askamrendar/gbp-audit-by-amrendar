import unittest,tempfile,pathlib,sys,json
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parents[1]))
from report import photo_count,rank_state,source_conflicts,export_portfolio
from audit import parse_photo_label

def record(code,position=None,valid=True):
 entries=[{'entity_id':'cid-other','position':1},{'entity_id':'cid-third','position':2},{'entity_id':'cid-fourth','position':3}]
 if position:entries.append({'entity_id':'cid-'+code,'position':position})
 return {'input':{'store_code':code,'keyword':'shoe store near me'},'started_at':'2026-10-22','profiles':[{'entity_id':'cid-'+code,'name':'Example '+code,'checks':[],'photo_count_observation':{'total':123,'observed_at':'2026-10-22'}}],'ranking':{'valid':valid,'results':entries,'method':'exploratory_maps_area_view'}}
class ReportTests(unittest.TestCase):
 def test_label_total(self):self.assertEqual(parse_photo_label('1,234 Photos')['total'],1234)
 def test_zero_is_observed(self):self.assertEqual(parse_photo_label('0 Photos')['total'],0)
 def test_media_not_mislabeled(self):self.assertEqual(parse_photo_label('101 Photos & videos')['label_kind'],'media')
 def test_non_count_label(self):self.assertIsNone(parse_photo_label('See photos'))
 def test_no_sample_as_total(self):self.assertEqual(photo_count({'gallery':{'items':[{}]*10}}),'Not exposed')
 def test_no_bad_comma_count(self):self.assertIsNone(parse_photo_label('1,,23 Photos'))
 def test_zero_total(self):self.assertEqual(photo_count({'photo_count_observation':{'total':0}}),'0')
 def test_failed_rank_not_zero_or_deep(self):self.assertEqual(rank_state(record('A',1,False))['label'],'Not measured')
 def test_censored_not_k_plus_one(self):self.assertEqual(rank_state(record('A'))['label'],'Not seen in 3 loaded results')
 def test_input_conflict_exact_values(self):
  data=record('A');data['input']['phone']='123';data['profiles'][0]['phone']='Phone: 456';c=source_conflicts(data)[0];self.assertEqual(c['observed'],'Phone: 456');self.assertEqual(c['source_value'],'123')
 def test_formatted_phone_not_false_conflict(self):
  data=record('A');data['input']['phone']='123-456';data['profiles'][0]['phone']='Phone: 123 456';self.assertFalse(source_conflicts(data))
 def test_sorted_bulk_and_merged_plan(self):
  from pptx import Presentation
  with tempfile.TemporaryDirectory() as tmp:
   file=pathlib.Path(tmp)/'bulk.pptx';export_portfolio([record('FOUR',4),record('UNKNOWN',valid=False),record('ONE',1),record('THREE',3)],file);r=Presentation(file)
   self.assertEqual(len(r.slides),13)
   titles=[s.shapes[0].text for s in r.slides]
   self.assertIn('(ONE)',titles[2]);self.assertIn('(THREE)',titles[4]);self.assertIn('(FOUR)',titles[6]);self.assertIn('(UNKNOWN)',titles[8]);self.assertEqual(titles[-2],'Profile work: setup and regular checks');self.assertEqual(titles[-1],'The next 30 days')
   for index in [3,5,7,9]:self.assertIn('Photos / media shown',' '.join(c.text for sh in r.slides[index].shapes if sh.has_table for row in sh.table.rows for c in row.cells))
 def test_different_keywords_cannot_share_rank_table(self):
  a,b=record('A'),record('B');b['input']['keyword']='other'
  with self.assertRaises(ValueError):export_portfolio([a,b],'unused.pptx')
 def test_equal_ranks_keep_order(self):self.assertEqual(rank_state(record('A',4))['sort'],rank_state(record('B',4))['sort'])
if __name__=='__main__':unittest.main()
