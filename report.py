"""Editable agency-style report; all conclusions retain collection limits."""
import pathlib,re
from collections import Counter
import json
from urllib.parse import urlparse

def website_facts(w):
 if not w or w.get('error'):return {'Destination':w.get('error','No inspected GBP-linked page') if w else 'No inspected GBP-linked page'}
 types=[]
 def walk(v):
  if isinstance(v,dict):
   t=v.get('@type',[]);types.extend(t if isinstance(t,list) else [t]);[walk(x) for x in v.values()]
  elif isinstance(v,list):[walk(x) for x in v]
 for block in w.get('schema',[]):walk(block.get('value'))
 frames=w.get('iframes',[]);social=[x for x in frames if re.search('instagram|facebook|tiktok|youtube',x.get('src',''),re.I)];maps=[x for x in frames if re.search('google.+maps',x.get('src',''),re.I)]
 return {'Destination':w.get('page_type','Unknown')+' · '+w.get('url',w.get('requested_url','')),'Title / local intent':w.get('title') or 'No title captured','Headings / layout':' | '.join(x.get('text','') for x in w.get('headings',[])[:3]) or 'No visible headings captured','Schema types':', '.join(dict.fromkeys(str(x) for x in types if x)) or 'No types parsed; confirm other markup formats','JSON-LD syntax':str(sum(x.get('syntax')=='invalid' for x in w.get('schema',[])))+' invalid blocks / '+str(len(w.get('schema',[])))+' inspected','Map embed':str(len(maps))+' iframe candidates; branch match needs review','Social embed':str(len(social))+' iframe candidates; rendered content needs review','Review embed':'Not established by iframe inventory; inspect rendered page/widgets','Content / conversion':'Branch address, hours, useful copy and CTA need comparison to approved facts'}

def form_values(p,keyword):
 body=p.get('body','');g=p.get('gallery',{});items=g.get('items',[]);w=p.get('website_audit',{});name=p.get('name') or 'Unknown'
 return [name,(p.get('address') or 'Address not exposed')+' / '+(p.get('phone') or 'Phone not exposed'), 'Displayed name: '+name+'; no keyword-only name additions', 'Verify locality wording against real signage',p.get('website') or ('No link · Add website observed' if p.get('website_link_state')=='missing' else 'Not exposed'),p.get('category') or 'Primary / additional categories not established', 'Expanded weekly schedule captured; verify special hours' if p.get('hours_evidence') else 'Full weekly schedule not captured', 'About evidence captured; inspect payment attributes' if p.get('about_evidence') else 'Payment methods not established',p.get('description') or 'Description not established',p.get('reviews_summary') or 'Review/reply period and denominator not measured', 'Post captured; inspect dated content and CTA' if p.get('post_evidence') else 'Published posts not established','Review actual captured GBP description/posts', 'Action-link availability and journey need review',w.get('page_type','No inspected destination')+'; branch identity needs review',p.get('entity_id') or 'Exact ID unresolved','Current Q&A availability / answers not established',photo_count(p)+' displayed; '+str(len(items))+' gallery items captured (sample)', '30/90-day owner uploads unknown; sample is not full history', ', '.join(sorted({x.get('image_type') for x in items if x.get('image_type')})) or 'Captured images need visual classification',str(sum(x.get('format')=='video' for x in items))+' video labels in sample; content review pending',website_facts(w).get('Map embed','No inspected linked-page map')]

FORM=['Business name','Address / phone','Keyword in name','Locality in name','GBP website link','Categories','Full opening hours','Payment attributes','Description','Review replies','Published posts','Copy / content type','Action links','Branch deep link','Profile identifier','Q&A availability','Gallery presence','Upload recency / quantity','Image types','Videos','Website map']

FORMAT_VERSION='2.0'

def photo_count(p):
 """A displayed total is never derived from the downloaded gallery sample."""
 observation=p.get('photo_count_observation') or {}
 if observation.get('total') is not None:return str(observation['total'])
 for key in ['displayed_photo_count','photo_count_displayed']:
  if p.get(key) is not None:return str(p[key])
 return 'Not exposed'

def rank_state(data):
 r=data.get('ranking') or {};t=(data.get('profiles') or [{}])[0]
 if data.get('identity_review_required') or not r.get('valid') or not t.get('entity_id'):return {'sort':(2,0),'label':'Not measured','position':None}
 for result in r.get('results',[]):
  if result.get('entity_id')==t['entity_id']:
   value=result.get('position')
   if isinstance(value,int) and value>0:return {'sort':(0,value),'label':'#'+str(value),'position':value}
 depth=len(r.get('results',[]))
 return {'sort':(1,0),'label':f'Not seen in {depth} loaded results','position':None}

def source_conflicts(data):
 p=(data.get('profiles') or [{}])[0];row=data['input'];result=[]
 for expected_key,actual_key,label in [('expected_address','address','Address'),('phone','phone','Phone'),('expected_hours','weekly_hours','Hours')]:
  expected=row.get(expected_key);actual=p.get(actual_key)
  if not expected or not actual:continue
  actual=' / '.join(map(str,actual)) if isinstance(actual,list) else str(actual)
  normalize=lambda text:re.sub(r'[^a-z0-9]','',re.sub(r'^(Address|Phone):\s*','',str(text),flags=re.I).casefold())
  if normalize(expected)!=normalize(actual):result.append({'field':label,'observed':actual,'source_value':str(expected),'source':'User-supplied input; operational truth unconfirmed','action':'Confirm the approved '+label.lower()+' with the store before changing either source.'})
 for item in p.get('conflicts',[]):
  if isinstance(item,dict) and item.get('observed') and item.get('source_value'):result.append(item)
 return result

def core_values(p):
 w=p.get('website_audit') or {};hours=p.get('weekly_hours') or p.get('live_hours')
 if isinstance(hours,list):hours=' / '.join(map(str,hours))
 if not hours:hours='Captured; verify full week' if p.get('hours_evidence') else 'Not established'
 site='No link shown' if p.get('website_link_state')=='missing' else (w.get('page_type') or 'Linked; branch to verify') if p.get('website') else 'Not exposed'
 rating=p.get('rating');reviews=p.get('review_count') or p.get('review_count_label')
 return [site,p.get('phone') or 'Not exposed',p.get('category') or 'Not exposed',hours,(str(rating) if rating is not None else 'Not exposed')+' / '+(str(reviews) if reviews is not None else 'not exposed'),photo_count(p)]

def first_action(data):
 p=data['profiles'][0];conflicts=source_conflicts(data)
 if conflicts:return conflicts[0].get('action','Confirm the conflicting source facts')
 if p.get('website_link_state')=='missing':return 'Add an approved store website destination'
 fails=[c for c in p.get('checks',[]) if c['status']=='Fail']
 if fails:return fails[0].get('next_input') or fails[0]['check']
 if p.get('blocker'):return 'Resolve collection access and confirm branch identity'
 return 'Confirm store facts and complete the evidence gaps'

def export_ppt(data,path):
 return export_portfolio([data],path)

def export_portfolio(audits,path):
 """Vicinus-style 2N+5 flow; called separately for compatible keyword/method groups."""
 if not audits:raise ValueError('No audit records to report')
 contexts={(d['input']['keyword'],d.get('ranking',{}).get('method','exploratory_maps_area_view')) for d in audits}
 if len(contexts)!=1:raise ValueError('Keep different keywords or ranking methods in separate portfolio reports')
 from pptx import Presentation
 from pptx.util import Inches,Pt
 from pptx.dml.color import RGBColor
 from PIL import Image
 ordered=sorted(enumerate(audits),key=lambda pair:(rank_state(pair[1])['sort'],pair[0]));audits=[d for _,d in ordered]
 r=Presentation();r.slide_width=Inches(13.333333);r.slide_height=Inches(7.5)
 brand=audits[0].get('brand') or {};ink=brand.get('color') or '#14213D';accent=brand.get('accent') or '#2F6FD6'
 if not re.fullmatch('#[a-fA-F0-9]{6}',ink):ink='#14213D'
 if not re.fullmatch('#[a-fA-F0-9]{6}',accent):accent='#2F6FD6'
 agency=brand.get('agency') or 'GBP Audit by Amrendar';client=brand.get('client') or audits[0]['profiles'][0].get('name') or 'Business audit'
 def rgb(c):return RGBColor.from_string(c.lstrip('#').upper())
 def tx(s,value,x,y,w,h,size=16,bold=False,color=None):
  sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.text_frame.word_wrap=True;sh.text_frame.text=str(value)
  for p in sh.text_frame.paragraphs:p.font.name='Calibri';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color or ink)
  return sh
 def fit_image(s,file,x,y,w,h):
  if not file or not pathlib.Path(file).is_file():tx(s,'Screenshot unavailable\nSee collection gap in audit.json',x,y,w,h,14,color='#6B7280');return
  try:
   with Image.open(file) as im:ratio=im.width/im.height
   width=min(w,h*ratio);height=width/ratio;s.shapes.add_picture(str(file),Inches(x+(w-width)/2),Inches(y+(h-height)/2),width=Inches(width),height=Inches(height))
  except (OSError,ValueError):tx(s,'Screenshot could not be opened',x,y,w,h,14)
 def slide(title,subtitle=''):
  s=r.slides.add_slide(r.slide_layouts[6]);tx(s,title,.6,.4,10.6,.65,25,True)
  if brand.get('logo_path'):fit_image(s,brand['logo_path'],11.65,.35,1.1,.38)
  else:tx(s,agency,11.15,.35,1.6,.4,9,False,accent)
  if subtitle:tx(s,subtitle,.6,1,12.1,.45,12,color='#6B7280')
  tx(s,str(len(r.slides)),12.4,7.12,.3,.2,9,color='#6B7280')
  return s
 def note(s,d):s.notes_slide.notes_text_frame.text=json.dumps(d,indent=2,ensure_ascii=False,default=str)
 def table(s,headers,rows,x=.6,y=1.65,w=12.1,h=4.3,size=12,widths=None):
  t=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h)).table
  if widths is None:widths=[2.1]+[(w-2.1)/(len(headers)-1)]*(len(headers)-1)
  for column,width in zip(t.columns,widths):column.width=Inches(width)
  for i,row in enumerate([headers]+rows):
   for j,value in enumerate(row):
    text=str(value);cell=t.cell(i,j);cell.text=text if len(text)<=130 else text[:127]+'…';cell.margin_left=cell.margin_right=Inches(.09);cell.margin_top=cell.margin_bottom=Inches(.04);cell.fill.solid();cell.fill.fore_color.rgb=rgb('#EEF1F6' if i==0 else '#FFFFFF')
    for p in cell.text_frame.paragraphs:p.font.name='Calibri';p.font.size=Pt(size);p.font.bold=i==0 or j==0;p.font.color.rgb=rgb(ink if i==0 or j==0 else '#222B38')
  note(s,{'full_table':[headers]+rows,'format_version':FORMAT_VERSION,'limits':'Unknown is not failure. Counts are not ranking requirements.'})
  return t
 keyword=audits[0]['input']['keyword'];s=slide(client)
 s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('#0D1D3A')
 for shape in s.shapes:
  if shape.has_text_frame:
   for paragraph in shape.text_frame.paragraphs:paragraph.font.color.rgb=rgb('#FFFFFF')
 tx(s,'Local Search Visibility Audit',.8,2.4,11.5,1,34,True,'#FFFFFF');tx(s,f'{len(audits)} store(s) · {keyword}',.8,3.55,11.5,.7,21,color='#FFFFFF')
 tx(s,'Public evidence report · '+audits[0]['started_at'][:10],.8,5.9,11.5,.5,13,color='#FFFFFF');note(s,{'observed_dates':[d['started_at'] for d in audits],'scope':'Public evidence; owner and visual checks may remain open','format_version':FORMAT_VERSION})
 # Paginate large bulk summaries; never squeeze 100 rows onto one slide.
 for offset in range(0,len(audits),8):
  group=audits[offset:offset+8];s=slide('Profile summary','Saved Maps order, not verified customer-location ranking. No missing position is invented.')
  rows=[[d['input']['store_code'],rank_state(d)['label'],core_values(d['profiles'][0])[0],first_action(d)] for d in group]
  table(s,['Store','Saved Maps position','Website link','First action'],rows,h=min(4.5,.65*(len(rows)+1)),size=11.5,widths=[1.5,2.4,2.5,5.7])
  tx(s,'Known positions first; not seen within the loaded depth and unmeasured views follow. Equal ranks keep input order.',.6,6.45,12,.45,11,color='#6B7280')
 for d in audits:
  p=d['profiles'][0];code=d['input']['store_code'];name=p.get('name') or code;names=[(e.get('name') or 'Unresolved') for e in d['profiles']];names[0]+=' ('+code+')';conflicts=source_conflicts(d)
  s=slide(name+' ('+code+'): what we found',keyword+' · '+rank_state(d)['label']+' · '+d['started_at'][:10])
  tx(s,'Business profile',.6,1.55,3.25,.3,12,True);tx(s,'Keyword results',3.95,1.55,3.25,.3,12,True)
  shot=(p.get('details_evidence') or {}).get('screenshot') or p.get('screenshot') or (p.get('evidence') or {}).get('screenshot');fit_image(s,shot,.6,1.95,3.25,3.55);fit_image(s,d.get('ranking',{}).get('screenshot'),3.95,1.95,3.25,3.55)
  rows=[]
  for c in conflicts[:3]:rows.extend([[c['field']+' — GBP',c['observed']],[c['field']+' — supplied source',c['source_value']]])
  if not rows:
   rows=[['Displayed address',p.get('address') or 'Not exposed'],['Website',core_values(p)[0]],['Collection',p.get('blocker') or 'Public profile evidence; operational truth unconfirmed']]
  table(s,['Where','What it says'],rows[:4],x=7.55,y=1.95,w=5.15,h=2.9,size=12,widths=[1.65,3.5])
  tx(s,'What to do',7.55,5.05,5.15,.35,16,True,accent);tx(s,first_action(d),7.55,5.48,5.15,.85,15)
  tx(s,'A source difference needs confirmation; it does not prove which source is correct.',.6,6.55,12,.35,10,color='#6B7280');note(s,{'source_conflicts':conflicts,'action':first_action(d),'profile_url':p.get('profile_url',p.get('url')),'profile_evidence':shot,'keyword_evidence':d.get('ranking',{}),'next_inputs':[c for c in p.get('checks',[]) if c['status'] in ['Fail','Partial','Unknown']],'form_comparison':d.get('form_comparison',[]),'website_audits':[{'entity':e.get('name'),'linked':e.get('website_audit'),'supplementary':e.get('supplementary_pages')} for e in d['profiles']]})
  s=slide(name+' ('+code+'): nearby stores','Same observed keyword cohort; primary categories only unless independently confirmed.')
  values=[core_values(e) for e in d['profiles']];fields=['Website destination','Phone','Main category','Opening hours','Rating / reviews','Photos / media shown']
  table(s,['Detail']+names,[[label]+[v[i] for v in values] for i,label in enumerate(fields)],h=4.2,size=11.5)
  tx(s,'Photo totals are displayed observations, not sample size or owner upload frequency. Capture times are retained in notes.',.6,6.2,12,.5,11,color='#6B7280');note(s,{'comparison':dict(zip(fields,[[v[i] for v in values] for i in range(6)])),'profile_observed_at':[e.get('observed_at') for e in d['profiles']],'photo_observations':[e.get('photo_count_observation') for e in d['profiles']],'websites':[website_facts(e.get('website_audit')) for e in d['profiles']],'checks':[(e.get('name'),e.get('checks',[])) for e in d['profiles']]})
 # Pilot checklist in the combined deck; individual store reports retain their own detail.
 d=audits[0];profiles=d['profiles'];names=[e.get('name') or 'Unresolved' for e in profiles];names[0]+=' ('+d['input']['store_code']+')'
 s=slide(d['input']['store_code']+': profile checklist vs competitors','Displayed totals, inspected samples and private history are separate measures.')
 def detailed(e):
  items=e.get('gallery',{}).get('items',[]);sample=f"{len(items)} items / {sum(x.get('format')=='video' for x in items)} video labels"
  attrs=e.get('attributes') or {}
  return [*core_values(e),sample,', '.join(sorted({x['image_type'] for x in items if x.get('image_type')})) or 'Visual review needed',e.get('reviews_summary') or 'Not measured','Captured; review content' if e.get('post_evidence') else 'Not exposed',attrs.get('payments','Not established'),'Link captured; verify' if e.get('whatsapp_links') else 'Not established']
 fields=['Website','Phone','Category','Hours','Rating / reviews','Photos / media shown','Gallery sample','Image types / relevance','Review replies','Posts','Payments','WhatsApp']
 vals=[detailed(e) for e in profiles];table(s,['Item']+names,[[f]+[v[i] for v in vals] for i,f in enumerate(fields)],y=1.5,h=5.15,size=10)
 note(s,{'pilot_branch':d['input']['store_code'],'form21_all_fields':[{ 'entity':e.get('name'),'values':dict(zip(FORM,form_values(e,keyword)))} for e in profiles],'website_schema_embeds':[website_facts(e.get('website_audit')) for e in profiles],'limits':'No automatic visual classification or full upload history inferred. Unexposed attributes remain unknown.'})
 s=slide('Profile work: setup and regular checks','Complete supported facts first, then maintain a useful customer journey.')
 table(s,['Area','Set up first','Keep doing'],[
 ['Store details','Confirm address, pin, floor, phone, hours and true categories/attributes.','Check weekly and set special hours before relevant dates.'],
 ['Website / products','Link the correct store page and use tracking tags. Verify eligible products.','Test links and update visit facts or stock-dependent claims.'],
 ['Description / questions','Write factual, naturally relevant text; put useful FAQs on the store page.','Answer real questions where supported; do not seed fake Q&A.'],
 ['WhatsApp','Add an approved link only where supported; assign staff and reply hours.','Monitor enquiries; do not promise unconfirmed stock.'],
 ['Posts / photos','Prepare genuine branch media and inspect unrelated images.','Start with useful weekly posts; refresh media when needed.'],
 ['Reviews','Create a neutral request and context-specific reply process.','Reply promptly to real feedback; no incentives or gating.'],
 ['Performance','Record keyword, location, method and owner-data access.','Compare equivalent observations, calls, directions and clicks.']],h=4.8,size=11.5,widths=[1.6,5.25,5.25])
 tx(s,'Suggested routine, not a Google ranking requirement. No uploads, edits or replies have been made.',.6,6.6,12,.3,10,color='#6B7280')
 s=slide('The next 30 days','Confirm facts, implement approved setup, start the routine and measure.')
 table(s,['When','Work','Done when'],[
 ['Days 1–7','Store teams confirm disputed facts and the correct contact/entrance. Assign owners.','One approved fact sheet per store; blocking inputs are identified.'],
 ['Days 8–14','Set approved website/contact links, factual descriptions and supported attributes.','Changes match approved facts; page, phone and enquiry journeys tested.'],
 ['Days 15–21','Begin useful posts, branch photos and review replies; staff WhatsApp where supported.','First content verified; test enquiry reaches the right team.'],
 ['Days 22–30','Repeat comparable saved views; establish a separate verified-origin baseline if available.','Dated report of completed fixes, enquiries and unresolved evidence.']],h=4.3,size=13,widths=[1.5,5.3,5.3])
 tx(s,'Live edits and publication need approval. No top-three deadline or ranking gain is guaranteed.',.6,6.3,12,.6,12,color='#6B7280')
 r.save(path)
 return str(path)
